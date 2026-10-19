"""Deterministic synthetic HR / recruitment datasets for the C817 labs.

Scenario: Orchid Bay Talent Partners - a FICTIONAL Singapore recruitment and HR
services firm running its Q3 2026 hiring drive: screening candidates for eight
client roles (12 openings). Every company, person, ID and value is synthetic and
generated from a fixed random seed, so the values quoted in each lab's "Test it"
step are reproducible.

One candidate pool runs through the whole course:
  * YTD   - 720 applications, 1 Jan - 26 Sep 2026 (Labs 7-10)
  * Q3    - the 240 of them received 1 Jul - 26 Sep 2026 (Labs 1-6)
  * RAW   - the Q3 export as it arrives from the applicant tracking system,
            with duplicates, spelling variants and other issues (Labs 1-3)
  * FINALISTS - 16 final-round candidates for four roles (Labs 11-12)

`build_lab_data(num, folder)` writes the lab's files into <folder>/data/.
`EXPECTED[num]` holds the values the lab text quotes for verification.
"""
import itertools
import os
import random
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COMPANY = "Orchid Bay Talent Partners"
FICTION_NOTE = (
    "All organisations, people, IDs and values in this workbook are FICTIONAL training data created for the "
    "Gemini for Google Sheet course (C817). Any resemblance to a real company or person is coincidental. "
    "Email addresses use the reserved example.com domain."
)
REPORT_DATE = date(2026, 9, 30)
OPEN_NOTE = (
    "Open in Google Sheets: upload this .xlsx to Google Drive, open it, then File > Save as Google Sheets. "
    "Gemini works best on a native Google Sheets file."
)

# Job_ID, Position, Client, Department, Min_Years, Required_Skills, Budget_Min, Budget_Max, Openings, Hiring_Manager, Recruiter
JOBS = [
    ("J101", "Data Analyst", "Larkspur Financial", "Finance & Analytics", 2, "SQL; Excel; Tableau", 5000, 6500, 2, "Melissa Tan", "Nur Aisyah"),
    ("J102", "HR Executive", "Orchid Bay (internal)", "Human Resources", 2, "HRIS; Recruitment; Employment Act", 3800, 4800, 1, "Grace Lim", "Meera Pillai"),
    ("J103", "Payroll Specialist", "Tembusu Logistics", "Finance & Analytics", 3, "Payroll; CPF; Excel", 4200, 5200, 1, "Hafiz Rahman", "Nur Aisyah"),
    ("J104", "Software Engineer", "Kaya Cloud Systems", "Technology", 3, "Python; JavaScript; Cloud", 6500, 8500, 2, "Arjun Menon", "Jonathan Lee"),
    ("J105", "Account Manager", "Seastar Retail Group", "Sales", 4, "B2B Sales; CRM; Negotiation", 5500, 7000, 1, "Sharon Koh", "Ryan Goh"),
    ("J106", "Customer Success Executive", "Kaya Cloud Systems", "Technology", 1, "CRM; Communication; SaaS", 3500, 4500, 2, "Arjun Menon", "Jonathan Lee"),
    ("J107", "Operations Supervisor", "Tembusu Logistics", "Operations", 5, "Warehouse; Team Leadership; Lean", 4500, 5800, 1, "Hafiz Rahman", "Meera Pillai"),
    ("J108", "Finance Associate", "Larkspur Financial", "Finance & Analytics", 1, "Accounting; Excel; Reconciliation", 3600, 4400, 2, "Melissa Tan", "Nur Aisyah"),
]
JOB = {j[0]: j for j in JOBS}
POS_JOB = {j[1]: j for j in JOBS}
POSITIONS = [j[1] for j in JOBS]
POS_W = [40, 28, 22, 38, 24, 34, 22, 32]          # Q3 applications per role (sums to 240)
RECRUITERS = ["Nur Aisyah", "Jonathan Lee", "Meera Pillai", "Ryan Goh"]

SOURCES = ["LinkedIn", "JobStreet", "MyCareersFuture", "Indeed", "Referral", "Company Website", "Career Fair"]
SOURCE_W = [26, 20, 16, 10, 9, 12, 7]
SOURCE_BONUS = {"LinkedIn": 0, "JobStreet": -1, "MyCareersFuture": -2, "Indeed": -2, "Referral": 7,
                "Company Website": 2, "Career Fair": -4}
SOURCE_COST = {"LinkedIn": 42, "JobStreet": 25, "MyCareersFuture": 5, "Indeed": 18, "Referral": 55,
               "Company Website": 8, "Career Fair": 35}  # S$ per application
EDUCATION = ["Diploma", "Bachelor's Degree", "Master's Degree", "Professional Certificate"]
EDU_W = [30, 48, 12, 10]
EXTRA_SKILLS = ["Excel", "Communication", "Project Management", "Power BI", "Customer Service", "Problem Solving",
                "Google Workspace", "Stakeholder Management", "Data Entry", "Presentation"]
STAGES = ["Applied", "Screened", "Shortlisted", "Interviewed", "Offered", "Hired", "Rejected", "Withdrawn"]
MONTH_N = {1: 70, 2: 62, 3: 78, 4: 84, 5: 88, 6: 98, 7: 86, 8: 82, 9: 72}   # applications per month, 2026

FIRST = ["Wei Ling", "Jun Hao", "Siti", "Arjun", "Priya", "Daniel", "Nurul", "Marcus", "Hui Min", "Farhan",
         "Kavitha", "Ethan", "Rachel", "Imran", "Mei Xin", "Vikram", "Chloe", "Haziq", "Sharmila", "Kenneth",
         "Aisha", "Ravi", "Joanne", "Zhi Wei", "Nadia", "Bryan", "Lakshmi", "Hidayah", "Gabriel", "Shawn",
         "Amanda", "Dinesh", "Jia Hui", "Faizal", "Natalie", "Suresh", "Kai Xuan", "Aaliyah", "Benjamin", "Divya",
         "Hannah", "Irfan", "Jasmine", "Karthik", "Liyana", "Matthew", "Nicole", "Omar", "Pei Shan", "Rohan",
         "Sabrina", "Terence", "Umairah", "Vanessa", "Xavier", "Yi Ting", "Zara", "Clement", "Felicia", "Hakim"]
LAST = ["Tan", "Lim", "Ng", "Wong", "Lee", "Chua", "Goh", "Teo", "Koh", "Ong", "Rahman", "Ismail", "Abdullah",
        "Kumar", "Nair", "Pillai", "Menon", "Singh", "Fernandez", "De Souza", "Hassan", "Yusof", "Chen", "Ho",
        "Low", "Yeo", "Sim", "Toh", "Ramasamy", "Krishnan", "Aziz", "Osman", "Leong", "Quek", "Seah", "Chong",
        "Pereira", "Iyer", "Salleh", "Foo"]

NOTICE_PHRASES = {
    0: ["is available immediately", "can start right away"],
    1: ["needs one week's notice", "has a 1-week notice period"],
    2: ["needs two weeks' notice", "has a 2-week notice period"],
    4: ["needs one month's notice", "has a 1-month notice period", "needs four weeks' notice"],
    8: ["needs two months' notice", "has an 8-week notice period"],
    12: ["needs three months' notice", "has a 3-month notice period"],
}
NOTICE_W = [0, 1, 2, 2, 4, 4, 4, 4, 8, 8, 12]
FEEDBACK = {
    "Positive": [
        "Strong {skill} skills and clear, structured answers. The panel recommends the next round.",
        "Impressive examples from past projects and a good team fit. Would hire.",
        "Confident, well prepared and solved the {skill} case quickly. Clear yes from the panel.",
    ],
    "Mixed": [
        "Good {skill} knowledge but struggled with the case question. Possible fit with coaching.",
        "Communicates well, though hands-on {skill} experience is limited. The panel was split.",
        "Solid background, but answers on {skill} lacked depth. Consider a second interview.",
    ],
    "Negative": [
        "Could not explain basic {skill} concepts and arrived late. Not recommended.",
        "Answers were vague and lacked examples. Do not proceed.",
        "Weak on {skill} and unclear about the role. The panel does not support an offer.",
    ],
}

HEADER_FILL = PatternFill("solid", fgColor="1A5E3A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1A5E3A")


# =========================================================================== the candidate pool
def _sentiment(score):
    return "Positive" if score >= 4.0 else ("Mixed" if score >= 3.0 else "Negative")


def _generate():
    rng = random.Random(817)
    names = [f"{f} {l}" for f in FIRST for l in LAST]
    rng.shuffle(names)
    rows, k = [], 0
    for m, n in MONTH_N.items():
        last_day = 26 if m == 9 else (date(2026, m + 1, 1) - timedelta(days=1)).day
        dates = sorted(date(2026, m, rng.randint(1, last_day)) for _ in range(n))
        if m >= 7:   # Q3: the exact role mix of the hiring drive
            start = {7: 0, 8: 86, 9: 168}[m]
            pos_list = _q3_positions()[start:start + n]
        else:
            pos_list = [rng.choices(POSITIONS, weights=POS_W)[0] for _ in range(n)]
        for d, pos in zip(dates, pos_list):
            k += 1
            rows.append(_candidate(rng, k, d, pos, names[k - 1]))
    return rows


def _q3_positions():
    pool = [p for p, w in zip(POSITIONS, POS_W) for _ in range(w)]
    random.Random(2026).shuffle(pool)
    return pool


def _candidate(rng, k, applied, pos, name):
    job = POS_JOB[pos]
    q3 = applied >= date(2026, 7, 1)
    job_id = job[0] if q3 else "J0" + job[0][2:]
    req = [s.strip() for s in job[5].split(";")]
    years = max(0, job[4] + rng.randint(-2, 8))
    n_req = rng.choices([0, 1, 2, 3], weights=[8, 25, 37, 30])[0]
    skills = rng.sample(req, n_req) + rng.sample([s for s in EXTRA_SKILLS if s not in req], rng.randint(1, 2))
    rng.shuffle(skills)
    source = rng.choices(SOURCES, weights=SOURCE_W)[0]
    score = 38 + 11 * n_req + min(max(years - job[4], -2), 5) * 3 + SOURCE_BONUS[source] + rng.randint(-8, 8)
    score = max(22, min(97, score))
    salary = int(round(rng.uniform(job[6] * 0.85, job[7] * 1.15) / 100.0)) * 100
    age = (REPORT_DATE - applied).days
    r = rng.random()

    def pick(table):
        acc = 0
        for stage, p in table:
            acc += p
            if r < acc:
                return stage
        return table[-1][0]

    if score >= 78:
        stage = pick([("Hired", .22), ("Offered", .08), ("Interviewed", .25), ("Shortlisted", .15), ("Rejected", .18), ("Withdrawn", .12)])
    elif score >= 65:
        stage = pick([("Hired", .05), ("Interviewed", .15), ("Shortlisted", .25), ("Screened", .20), ("Rejected", .30), ("Withdrawn", .05)])
    else:
        stage = pick([("Screened", .10), ("Applied", .10), ("Rejected", .80)])
    if age > 60:  # older applications are resolved
        stage = {"Applied": "Rejected", "Screened": "Rejected", "Shortlisted": "Rejected" if r < .7 else "Withdrawn",
                 "Interviewed": "Rejected" if r < .75 else "Withdrawn",
                 "Offered": "Hired" if r < .8 else "Withdrawn"}.get(stage, stage)
    if age < 14 and stage in ("Interviewed", "Offered", "Hired"):
        stage = "Shortlisted"
    if age < 6 and stage not in ("Rejected",):
        stage = "Applied" if r < .5 else "Screened"
    interview, idate, hdate = None, None, None
    if stage in ("Interviewed", "Offered", "Hired"):
        interview = round(min(5.0, max(1.6, 1 + score / 25 + rng.gauss(0, 0.55))), 1)
        idate = min(applied + timedelta(days=rng.randint(7, 16)), REPORT_DATE)
    if stage == "Hired":
        hdate = applied + timedelta(days=rng.randint(20, 46))
        if hdate > REPORT_DATE:
            stage, hdate = "Offered", None
    first, last = name.split(" ", 1) if name.count(" ") == 1 else (name.rsplit(" ", 1)[0], name.rsplit(" ", 1)[1])
    email = f"{first}.{last}".lower().replace(" ", "") + f"{k % 90 + 10}@example.com"
    notice = rng.choice(NOTICE_W)
    edu = rng.choices(EDUCATION, weights=EDU_W)[0]
    domain = {"Data Analyst": "data and reporting", "HR Executive": "HR operations", "Payroll Specialist": "payroll",
              "Software Engineer": "software development", "Account Manager": "B2B account management",
              "Customer Success Executive": "customer support", "Operations Supervisor": "warehouse operations",
              "Finance Associate": "finance and accounts"}[pos]
    summary = (f"{years} year{'s' if years != 1 else ''} in {domain}; skilled in {', '.join(skills)}. "
               f"Holds a {edu.lower()}. {rng.choice(['Candidate', 'Applicant'])} {rng.choice(NOTICE_PHRASES[notice])}.")
    fb_score = interview if interview is not None else (4.3 if score >= 85 else 3.5 if score >= 72 else 2.6)
    sentiment = _sentiment(fb_score)
    feedback = rng.choice(FEEDBACK[sentiment]).format(skill=rng.choice(req))
    return dict(
        Candidate_ID=f"CAN-{26000 + k}", Full_Name=name, Email=email, Applied_Date=applied,
        Month=f"2026-{applied.month:02d}", Job_ID=job_id, Position=pos, Source=source, Years_Experience=years,
        Highest_Education=edu, Key_Skills="; ".join(skills), Expected_Salary=salary, Notice_Weeks=notice,
        Screening_Score=score, Interview_Score=interview, Stage=stage, Recruiter=job[10], Interview_Date=idate,
        Hire_Date=hdate, Days_To_Hire=(hdate - applied).days if hdate else None, Summary=summary,
        Feedback=feedback, Sentiment=sentiment, Skills_List=skills, Req_Met=n_req, Q3=q3,
    )


YTD = _generate()
Q3 = [r for r in YTD if r["Q3"]]
assert len(YTD) == 720 and len(Q3) == 240


# =========================================================================== workbook helpers
def _write_rows(ws, headers, rows, formats=None, widths=None, start_row=1, freeze=True):
    formats = formats or {}
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3F1E8")
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = formats.get(headers[c - 1])
            if fmt:
                cell.number_format = fmt
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
    for c, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or max(11, min(30, len(h) + 4))
        ws.column_dimensions[get_column_letter(c)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _about_sheet(wb, title, purpose, dictionary):
    ws = wb.create_sheet("About_This_Data", 0)
    ws["A1"] = f"{COMPANY} — {title}"
    ws["A1"].font = TITLE_FONT
    for cell, text in (("A2", FICTION_NOTE), ("A3", purpose), ("A4", OPEN_NOTE)):
        ws[cell] = text
        ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(f"{cell}:C{cell[1:]}")
        ws.row_dimensions[int(cell[1:])].height = 46
    ws["A6"], ws["B6"], ws["C6"] = "Sheet / Column", "Type", "Meaning"
    for c in ("A6", "B6", "C6"):
        ws[c].fill = HEADER_FILL
        ws[c].font = HEADER_FONT
    for i, (col, typ, meaning) in enumerate(dictionary, 7):
        ws.cell(row=i, column=1, value=col)
        ws.cell(row=i, column=2, value=typ)
        ws.cell(row=i, column=3, value=meaning).alignment = Alignment(wrap_text=True)
        if len(meaning) > 85:
            ws.row_dimensions[i].height = 15 * -(-len(meaning) // 85)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90


def _prompt_log_sheet(wb):
    ws = wb.create_sheet("Prompt_Log")
    headers = ["Step", "Where (Ask Gemini / Insert / Action / =AI / Gem)", "Prompt", "What Gemini did",
               "How I verified", "Decision"]
    _write_rows(ws, headers, [], widths={"Prompt": 60, "What Gemini did": 40, "How I verified": 40, "Decision": 14})
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).fill = HEADER_FILL
        ws.cell(row=1, column=c).font = HEADER_FONT
    dv = DataValidation(type="list", formula1='"Keep,Revise,Undo"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F2:F60")


def _save(wb, folder, name):
    out = os.path.join(folder, "data")
    os.makedirs(out, exist_ok=True)
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        del wb["Sheet"]
    path = os.path.join(out, name)
    wb.save(path)
    return path


def _jobs_sheet(wb, jobs=JOBS, name="Jobs"):
    ws = wb.create_sheet(name)
    headers = ["Job_ID", "Position", "Client", "Department", "Min_Years", "Required_Skills", "Budget_Min",
               "Budget_Max", "Openings", "Hiring_Manager", "Recruiter", "Posted_Date"]
    rows = [list(j[:11]) + [date(2026, 6, 24) if j[0].startswith("J1") else date(2025, 12, 15)] for j in jobs]
    _write_rows(ws, headers, rows, formats={"Budget_Min": '"S$"#,##0', "Budget_Max": '"S$"#,##0'},
                widths={"Required_Skills": 34, "Client": 22, "Position": 26})
    return ws


_JOB_DICT = [
    ("Jobs (sheet)", "—", "The eight client roles in the Q3 2026 hiring drive - one row per job requisition."),
    ("Job_ID", "Text", "Requisition ID, e.g. J101."),
    ("Min_Years", "Number", "Minimum years of relevant experience for the role."),
    ("Required_Skills", "Text", "Must-have skills, separated by semicolons."),
    ("Budget_Min / Budget_Max", "Number", "Monthly salary band agreed with the client, in S$."),
    ("Openings", "Number", "How many people the client wants to hire for this role."),
]

APP_COLS = ["Candidate_ID", "Full_Name", "Email", "Applied_Date", "Job_ID", "Position", "Source", "Years_Experience",
            "Highest_Education", "Key_Skills", "Expected_Salary", "Notice_Weeks", "Screening_Score",
            "Interview_Score", "Stage", "Recruiter"]
_APP_DICT = [
    ("Applications (sheet)", "—", "One row per candidate application received in the Q3 2026 hiring drive."),
    ("Candidate_ID", "Text", "Unique application ID from the applicant tracking system (ATS)."),
    ("Full_Name / Email", "Text", "Fictional candidate name and an example.com email address."),
    ("Applied_Date", "Date", "Date the application was received (yyyy-mm-dd)."),
    ("Job_ID / Position", "Text", "The requisition and role applied for - see the Jobs sheet."),
    ("Source", "Text", "Where the candidate came from: LinkedIn, JobStreet, MyCareersFuture, Indeed, Referral, "
                       "Company Website or Career Fair."),
    ("Years_Experience", "Number", "Years of relevant experience stated by the candidate."),
    ("Key_Skills", "Text", "Skills listed on the application, separated by semicolons."),
    ("Expected_Salary", "Number", "Expected monthly salary in S$."),
    ("Notice_Weeks", "Number", "Notice period in weeks."),
    ("Screening_Score", "Number", "ATS screening score, 0-100 (skills, experience and fit)."),
    ("Interview_Score", "Number", "Panel interview score 1.0-5.0; blank if not interviewed yet."),
    ("Stage", "Text", "Applied, Screened, Shortlisted, Interviewed, Offered, Hired, Rejected or Withdrawn."),
    ("Recruiter", "Text", "Orchid Bay recruiter who owns the requisition."),
]


def _app_row(r):
    return [r[c] for c in APP_COLS]


# =========================================================================== LABS 1-3: the raw export
def _raw():
    """The Q3 export with realistic data-quality issues (fixed positions, so the counts are exact)."""
    rng = random.Random(4242)
    rows = [_app_row(r) for r in Q3]
    ix = {c: i for i, c in enumerate(APP_COLS)}
    li = [i for i, r in enumerate(rows) if r[ix["Source"]] == "LinkedIn"]
    js = [i for i, r in enumerate(rows) if r[ix["Source"]] == "JobStreet"]
    picks_li, picks_js = rng.sample(li, 5), rng.sample(js, 3)
    for i in picks_li[:3]:
        rows[i][ix["Source"]] = "Linked In"
    for i in picks_li[3:]:
        rows[i][ix["Source"]] = "linkedin"
    for i in picks_js:
        rows[i][ix["Source"]] = "JobStreet "
    others = [i for i in range(len(rows)) if i not in picks_li + picks_js]
    sample = rng.sample(others, 4 + 2 + 5 + 1 + 4 + 3)
    miss, bad, sal, neg, caps, space = (sample[:4], sample[4:6], sample[6:11], sample[11:12], sample[12:16],
                                        sample[16:19])
    for i in miss:
        rows[i][ix["Email"]] = None
    for i in bad:
        rows[i][ix["Email"]] = rows[i][ix["Email"]].replace("@", ".")
    for i in sal:
        rows[i][ix["Expected_Salary"]] = f"S${rows[i][ix['Expected_Salary']]:,}"
    neg_true = rows[neg[0]][ix["Years_Experience"]]
    rows[neg[0]][ix["Years_Experience"]] = -neg_true if neg_true else -3
    for i in caps:
        rows[i][ix["Full_Name"]] = rows[i][ix["Full_Name"]].upper()
    for i in space:
        rows[i][ix["Full_Name"]] = rows[i][ix["Full_Name"]].replace(" ", "  ", 1) + " "
    notes_pool = ["Called twice - no answer", "Prefers hybrid work", "Asked about bonus", "Re-applied from 2025",
                  "Check visa status with client", "Referred by ex-colleague", "Strong portfolio"]
    notes = [rng.choice(notes_pool) if rng.random() < 0.18 else None for _ in rows]
    # 3 exact duplicate rows (the ATS exported them twice): copies sit near the end
    dup_src = sorted(rng.sample([i for i in others if i not in sample][:120], 3))
    final, final_notes = list(rows), list(notes)
    for j, i in enumerate(dup_src):
        pos = len(final) - (5 + 9 * j)
        final.insert(pos, list(rows[i]))
        final_notes.insert(pos, notes[i])
    issues = dict(
        rows=len(final), unique=len(rows), dup_ids=[rows[i][0] for i in dup_src],
        linked_in=3, linkedin_lower=2, jobstreet_space=3, missing_email=4, bad_email=2,
        bad_email_ids=[rows[i][0] for i in bad], text_salary=5, neg_years=1, neg_id=rows[neg[0]][0],
        neg_true=neg_true if neg_true else 3, caps_names=4, space_names=3,
    )
    return final, final_notes, issues


RAW_ROWS, RAW_NOTES, RAW_ISSUES = _raw()


def _raw_workbook(title, purpose, with_notes=False, extra=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    headers = APP_COLS + (["Internal_Notes"] if with_notes else [])
    rows = [r + ([n] if with_notes else []) for r, n in zip(RAW_ROWS, RAW_NOTES)]
    _write_rows(ws, headers, rows, formats={"Expected_Salary": "#,##0"},
                widths={"Full_Name": 24, "Email": 32, "Key_Skills": 36, "Position": 26, "Internal_Notes": 28})
    _jobs_sheet(wb)
    if extra:
        extra(wb)
    _prompt_log_sheet(wb)
    _about_sheet(wb, title, purpose, _APP_DICT + _JOB_DICT +
                 ([("Internal_Notes", "Text", "Free-text recruiter notes - not needed for screening.")] if with_notes else []))
    return wb


def _q3_stats(rows):
    pos = Counter(r["Position"] for r in rows)
    src = Counter(r["Source"] for r in rows)
    return pos, src


def _lab1():
    pos = Counter(r[5] for r in RAW_ROWS)
    top = pos.most_common(1)[0]
    pos_clean, _ = _q3_stats(Q3)
    dates = [r[3] for r in RAW_ROWS]
    sal = [r[10] for r in RAW_ROWS if isinstance(r[10], int)]
    return dict(rows=len(RAW_ROWS), cols=len(APP_COLS), first=min(dates), last=max(dates), positions=8,
                pos=dict(pos), openings={j[1]: j[8] for j in JOBS},
                top_pos=top[0], top_pos_n=top[1], top_pos_clean=pos_clean[top[0]],
                se=pos["Software Engineer"], recruiters=4,
                stage=Counter(r[14] for r in RAW_ROWS), avg_salary_numeric=round(sum(sal) / len(sal)))


def build_lab1(folder):
    wb = _raw_workbook("Q3 2026 Applications Export",
                       "Raw export of the Q3 2026 hiring drive from the applicant tracking system (ATS), exactly as "
                       "it arrived. Lab 1 explores it with Gemini; it has NOT been checked or cleaned yet.")
    return [_save(wb, folder, "OrchidBay_Q3_Applications_Export.xlsx")]


def _lab2():
    return dict(RAW_ISSUES)


def build_lab2(folder):
    def checks(wb):
        ws = wb.create_sheet("Checks")
        ws["A1"], ws["B1"], ws["C1"] = "Check", "Result (your formula)", "Expected"
        for c in ("A1", "B1", "C1"):
            ws[c].font = Font(bold=True)
        e = RAW_ISSUES
        for i, (lab, val) in enumerate([("Total rows", e["rows"]), ("Unique IDs", e["unique"]),
                                         ("Blank emails", e["missing_email"]), ("Emails without @", e["bad_email"]),
                                         ("Text salaries", e["text_salary"]), ("Negative experience", e["neg_years"])], 2):
            ws.cell(row=i, column=1, value=lab)
            ws.cell(row=i, column=3, value=val)
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 14
    wb = _raw_workbook("Q3 2026 Applications - Quality Check",
                       "The same raw Q3 export used in Lab 1. In Lab 2 you check its data quality with Gemini before "
                       "anyone analyses it. Do not fix anything yet - Lab 3 does the clean-up.", extra=checks)
    return [_save(wb, folder, "OrchidBay_Q3_Applications_QualityCheck.xlsx")]


def _lab3():
    clean_src = Counter(r["Source"] for r in Q3)
    se_ref = sum(1 for r in Q3 if r["Position"] == "Software Engineer" and r["Source"] == "Referral")
    newest = max(r["Applied_Date"] for r in Q3)
    return dict(rows_after=240, removed=3, sources=7, hired=sum(1 for r in Q3 if r["Stage"] == "Hired"), linkedin=clean_src["LinkedIn"], jobstreet=clean_src["JobStreet"],
                se_referral=se_ref, newest=newest, cols_after=len(APP_COLS), neg_id=RAW_ISSUES["neg_id"],
                neg_true=RAW_ISSUES["neg_true"])


def build_lab3(folder):
    wb = _raw_workbook("Q3 2026 Applications - Data Prep",
                       "The raw Q3 export again, with the recruiters' Internal_Notes column. In Lab 3 you clean and "
                       "prepare it with Gemini actions. Recruiter note: the candidate with a negative "
                       f"Years_Experience ({RAW_ISSUES['neg_id']}) confirmed by phone that the correct value is "
                       f"{RAW_ISSUES['neg_true']}.", with_notes=True)
    return [_save(wb, folder, "OrchidBay_Q3_Applications_Prep.xlsx")]


# =========================================================================== LAB 4: formulas
def _screen(r):
    job = JOB[r["Job_ID"]]
    meets = r["Years_Experience"] >= job[4]
    if r["Screening_Score"] >= 70 and meets:
        return "Strong"
    return "Possible" if r["Screening_Score"] >= 60 else "Not suitable"


def _lab4():
    res = Counter(_screen(r) for r in Q3)
    above = sum(1 for r in Q3 if r["Expected_Salary"] > JOB[r["Job_ID"]][7])
    da_strong = sum(1 for r in Q3 if r["Position"] == "Data Analyst" and _screen(r) == "Strong")
    first = Q3[0]
    return dict(rows=240, strong=res["Strong"], possible=res["Possible"], not_suitable=res["Not suitable"],
                above=above, within=240 - above, da_strong=da_strong,
                first_id=first["Candidate_ID"], first_job=first["Job_ID"], first_pos=first["Position"],
                first_min=JOB[first["Job_ID"]][4], first_budget=JOB[first["Job_ID"]][7],
                first_result=_screen(first), first_score=first["Screening_Score"], first_years=first["Years_Experience"],
                first_salary=first["Expected_Salary"],
                strong75=sum(1 for r in Q3 if r["Screening_Score"] >= 75 and r["Years_Experience"] >= JOB[r["Job_ID"]][4]),
                strong_by_pos={p: sum(1 for r in Q3 if r["Position"] == p and _screen(r) == "Strong") for p in POSITIONS})


def build_lab4(folder):
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    headers = APP_COLS + ["Budget_Max"]
    rows = []
    for i, r in enumerate(Q3, 2):
        rows.append(_app_row(r) + [f"=VLOOKUP(E{i},Jobs!$A$2:$H$9,9,FALSE)"])
    _write_rows(ws, headers, rows, formats={"Expected_Salary": '"S$"#,##0', "Budget_Max": '"S$"#,##0'},
                widths={"Full_Name": 24, "Email": 32, "Key_Skills": 36, "Position": 26})
    _jobs_sheet(wb)
    st = wb.create_sheet("Settings")
    st["A1"], st["B1"] = "Report date", REPORT_DATE
    st["B1"].number_format = "yyyy-mm-dd"
    st["A2"], st["B2"] = "Pass mark (Screening_Score)", 70
    st["A3"], st["B3"] = "Possible mark (Screening_Score)", 60
    st.column_dimensions["A"].width = 34
    ck = wb.create_sheet("Checks")
    ck["A1"], ck["B1"] = "Check", "Result"
    ck["A1"].font = ck["B1"].font = Font(bold=True)
    ck.column_dimensions["A"].width = 46
    _prompt_log_sheet(wb)
    _about_sheet(wb, "Q3 2026 Applications - Screening Formulas",
                 "The cleaned Q3 applications (240 rows - the result of Lab 3). In Lab 4 you add screening formulas "
                 "with Gemini. Column Q (Budget_Max) already holds a formula with an error for you to fix with Gemini.",
                 _APP_DICT + [("Budget_Max", "Formula", "Should look up the role's Budget_Max from Jobs - it currently "
                                                         "shows an error.")] + _JOB_DICT +
                 [("Settings (sheet)", "—", "Report date, pass mark (70) and possible mark (60) used by the formulas.")])
    return [_save(wb, folder, "OrchidBay_Screening_Formulas.xlsx")]


# =========================================================================== LAB 5: tracker
SHORT_COLS = ["Candidate_ID", "Full_Name", "Position", "Source", "Screening_Score", "Interview_Score",
              "Expected_Salary", "Budget_Max", "Notice_Weeks", "Stage", "Recruiter"]
ACTIVE = ("Shortlisted", "Interviewed", "Offered")


def _shortlist():
    return [r for r in Q3 if r["Stage"] in ACTIVE]


def _interview_notes():
    rng = random.Random(55)
    cands = [r for r in Q3 if r["Stage"] == "Shortlisted"][:12]
    slots = [("Mon 28 Sep", "10:00"), ("Mon 28 Sep", "14:30"), ("Tue 29 Sep", "09:30"), ("Tue 29 Sep", "11:00"),
             ("Tue 29 Sep", "16:00"), ("Wed 30 Sep", "10:00"), ("Wed 30 Sep", "15:00"), ("Thu 1 Oct", "09:30"),
             ("Thu 1 Oct", "11:30"), ("Thu 1 Oct", "14:00"), ("Fri 2 Oct", "10:00"), ("Fri 2 Oct", "15:30")]
    status = ["Completed", "Completed", "No-show", "Completed", "Rescheduled", "Completed", "Completed",
              "Scheduled", "Rescheduled", "Scheduled", "Scheduled", "No-show"]
    phrase = {"Completed": ["done - panel to send scores", "interview done, notes in ATS", "completed, went well"],
              "No-show": ["no-show, tried calling twice", "candidate did not turn up"],
              "Rescheduled": ["candidate asked to move it to next week", "moved - hiring manager on leave"],
              "Scheduled": ["confirmed by email", "invite sent and accepted"]}
    lines = []
    for (day, t), r, s in zip(slots, cands, status):
        job = JOB[r["Job_ID"]]
        fmt = rng.choice(["video call", "on-site at client office", "phone"])
        lines.append(f"{day} {t} - {r['Full_Name']} ({r['Candidate_ID']}) for {r['Position']}. Panel: {job[10]} + "
                     f"{job[9]}. {fmt}. {rng.choice(phrase[s])}.")
    return lines, Counter(status), cands


def _lab5():
    sl = _shortlist()
    _lines, st, _c = _interview_notes()
    return dict(notes=12, status=dict(st), shortlist=len(sl),
                hi85=sum(1 for r in sl if r["Screening_Score"] >= 85),
                hi80=sum(1 for r in sl if r["Screening_Score"] >= 80),
                above=sum(1 for r in sl if r["Expected_Salary"] > JOB[r["Job_ID"]][7]),
                long_notice=sum(1 for r in sl if r["Notice_Weeks"] > 8),
                stages=dict(Counter(r["Stage"] for r in sl)))


def build_lab5(folder):
    wb = Workbook()
    ws = wb.active
    ws.title = "Interview_Notes"
    lines, _st, _c = _interview_notes()
    ws["A1"] = "Interview notes - week of 28 Sep 2026 (typed by recruiters during the week)"
    ws["A1"].font = Font(bold=True)
    for i, line in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 150
    sh = wb.create_sheet("Shortlist")
    rows = [[r["Candidate_ID"], r["Full_Name"], r["Position"], r["Source"], r["Screening_Score"],
             r["Interview_Score"], r["Expected_Salary"], JOB[r["Job_ID"]][7], r["Notice_Weeks"], r["Stage"],
             r["Recruiter"]] for r in _shortlist()]
    _write_rows(sh, SHORT_COLS, rows, formats={"Expected_Salary": '"S$"#,##0', "Budget_Max": '"S$"#,##0'},
                widths={"Full_Name": 24, "Position": 26})
    _prompt_log_sheet(wb)
    _about_sheet(wb, "Interview Tracker",
                 "This week's interview notes (unstructured text) and the active shortlist of the Q3 2026 hiring drive "
                 "(candidates who are Shortlisted, Interviewed or Offered). In Lab 5 you turn the notes into a colour-"
                 "coded tracker with dropdowns and checkboxes, and flag the shortlist with conditional formatting.",
                 [("Interview_Notes (sheet)", "Text", "One line per interview, as typed by recruiters."),
                  ("Shortlist (sheet)", "—", "One row per active candidate."),
                  ("Budget_Max", "Number", "Top of the client's monthly salary band for the role, S$."),
                  ("Notice_Weeks", "Number", "Notice period in weeks.")] + _APP_DICT[10:14])
    return [_save(wb, folder, "OrchidBay_Interview_Tracker.xlsx")]


# =========================================================================== LAB 6: =AI function
def _profiles():
    return [r for r in Q3 if r["Stage"] in ("Shortlisted", "Interviewed", "Offered", "Hired")][:30]


def _lab6():
    p = _profiles()
    sent = Counter(r["Sentiment"] for r in p)
    notice = Counter(r["Notice_Weeks"] for r in p)
    return dict(rows=len(p), positive=sent["Positive"], mixed=sent["Mixed"], negative=sent["Negative"],
                notice=dict(sorted(notice.items())), first_id=p[0]["Candidate_ID"], first_notice=p[0]["Notice_Weeks"],
                first_summary=p[0]["Summary"], first_sent=p[0]["Sentiment"], first_feedback=p[0]["Feedback"],
                long_notice=sum(1 for r in p if r["Notice_Weeks"] >= 8))


def build_lab6(folder):
    wb = Workbook()
    ws = wb.active
    ws.title = "Profiles"
    headers = ["Candidate_ID", "Full_Name", "Position", "Candidate_Summary", "Interview_Feedback",
               "Notice_Weeks_AI", "Feedback_Sentiment_AI", "Manager_Summary_AI"]
    rows = [[r["Candidate_ID"], r["Full_Name"], r["Position"], r["Summary"], r["Feedback"], None, None, None]
            for r in _profiles()]
    _write_rows(ws, headers, rows, widths={"Candidate_Summary": 60, "Interview_Feedback": 55, "Full_Name": 22,
                                            "Position": 24, "Manager_Summary_AI": 50, "Feedback_Sentiment_AI": 22})
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=4, max_col=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ak = wb.create_sheet("Answer_Key")
    _write_rows(ak, ["Candidate_ID", "Notice_Weeks", "Feedback_Sentiment"],
                [[r["Candidate_ID"], r["Notice_Weeks"], r["Sentiment"]] for r in _profiles()])
    ck = wb.create_sheet("Checks")
    ck["A1"], ck["B1"] = "Check", "Result"
    ck["A1"].font = ck["B1"].font = Font(bold=True)
    ck.column_dimensions["A"].width = 50
    _prompt_log_sheet(wb)
    _about_sheet(wb, "Candidate Profiles",
                 "30 candidates from the active Q3 2026 shortlist, with the free-text summary from their application "
                 "and the panel's interview or phone-screen feedback. In Lab 6 you extract and categorise this text "
                 "with the AI function (=AI) and Fill with Gemini. Answer_Key holds the recorded values so you can "
                 "measure how accurate the AI output is - never put it in an =AI prompt; compare against it only after you have generated your answers.",
                 [("Candidate_Summary", "Text", "Free text: experience, skills, education and notice period."),
                  ("Interview_Feedback", "Text", "Free text from the interview panel or phone screen."),
                  ("Notice_Weeks_AI", "Empty", "You fill this with =AI: notice period in weeks."),
                  ("Feedback_Sentiment_AI", "Empty", "You fill this with =AI: Positive, Mixed or Negative."),
                  ("Manager_Summary_AI", "Empty", "You fill this with =AI: one-sentence summary for the hiring manager."),
                  ("Answer_Key (sheet)", "—", "Recorded notice weeks and feedback sentiment, for checking only.")])
    return [_save(wb, folder, "OrchidBay_Candidate_Profiles.xlsx")]


# =========================================================================== LABS 7-10: 2026 YTD
YTD_COLS = ["Candidate_ID", "Full_Name", "Applied_Date", "Month", "Job_ID", "Position", "Source", "Years_Experience",
            "Highest_Education", "Expected_Salary", "Screening_Score", "Interview_Score", "Stage", "Recruiter",
            "Hire_Date", "Days_To_Hire"]
_YTD_DICT = [
    ("Applications (sheet)", "—", "One row per application received 1 Jan - 26 Sep 2026 (720 rows). Jul-Sep rows "
                                  "are the Q3 hiring drive you cleaned in Labs 1-3."),
    ("Month", "Text", "Application month, yyyy-mm (2026-01 to 2026-09)."),
    ("Interview_Score", "Number", "Panel interview score 1.0-5.0; blank if never interviewed."),
    ("Stage", "Text", "Current stage: Applied, Screened, Shortlisted, Interviewed, Offered, Hired, Rejected, Withdrawn."),
    ("Hire_Date / Days_To_Hire", "Date / Number", "For hired candidates: start-of-offer date and days from "
                                                   "application to hire."),
    ("Source_Costs (sheet)", "—", "Advertising / referral cost per application for each source, S$."),
    ("Data minimisation", "Note", "Emails, skills and notice periods were removed - analysis does not need them."),
    ("Jobs (sheet)", "—", "J101-J108 are the Q3 2026 requisitions; J001-J008 are the H1 2026 requisitions for the same "
                          "eight roles. Count Openings for one half only."),
]
H1_JOBS = [("J0" + j[0][2:],) + j[1:] for j in JOBS]


def _ytd_rows():
    return [[r[c] for c in YTD_COLS] for r in YTD]


def _ytd_workbook(title, purpose, extra=None, costs=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _write_rows(ws, YTD_COLS, _ytd_rows(), formats={"Expected_Salary": '"S$"#,##0'},
                widths={"Full_Name": 24, "Position": 26})
    _jobs_sheet(wb, JOBS + H1_JOBS)
    if costs:
        sc = wb.create_sheet("Source_Costs")
        _write_rows(sc, ["Source", "Cost_Per_Application"], [[s, SOURCE_COST[s]] for s in SOURCES],
                    formats={"Cost_Per_Application": '"S$"#,##0'}, widths={"Cost_Per_Application": 22})
    if extra:
        extra(wb)
    _prompt_log_sheet(wb)
    _about_sheet(wb, title, purpose, _YTD_DICT + _JOB_DICT[:1])
    return wb


def _src_stats():
    out = {}
    for s in SOURCES:
        rows = [r for r in YTD if r["Source"] == s]
        hires = sum(1 for r in rows if r["Stage"] == "Hired")
        cost = len(rows) * SOURCE_COST[s]
        out[s] = dict(apps=len(rows), hires=hires, rate=round(hires / len(rows) * 100, 1),
                      cost=cost, cph=round(cost / hires) if hires else None)
    return out


def _lab7():
    by_month = Counter(r["Month"] for r in YTD)
    hires = [r for r in YTD if r["Stage"] == "Hired"]
    hires_m = Counter(r["Month"] for r in hires)
    src = _src_stats()
    best = max(src, key=lambda s: src[s]["rate"])
    most = max(src, key=lambda s: src[s]["hires"])
    dth = {p: round(sum(r["Days_To_Hire"] for r in hires if r["Position"] == p) /
                    max(1, sum(1 for r in hires if r["Position"] == p)), 1) for p in POSITIONS}
    slow = max(dth, key=dth.get)
    stage = Counter(r["Stage"] for r in YTD)
    return dict(rows=720, by_month=dict(sorted(by_month.items())), peak=max(by_month, key=by_month.get),
                hires=len(hires), hires_by_month=dict(sorted(hires_m.items())), src=src, best_rate=best,
                most_hires=most, dth=dth, slowest=slow, avg_dth=round(sum(r["Days_To_Hire"] for r in hires) / len(hires), 1),
                stage=dict(stage), li_hired=sum(1 for r in YTD if r["Source"] == "LinkedIn" and r["Stage"] == "Hired"),
                li_rejected=sum(1 for r in YTD if r["Source"] == "LinkedIn" and r["Stage"] == "Rejected"))


def build_lab7(folder):
    wb = _ytd_workbook("Hiring 2026 Year to Date - Analysis",
                       "All 2026 applications to date across the eight roles. In Lab 7 you ask Gemini for analysis "
                       "and insights, build pivot tables and check every number.")
    return [_save(wb, folder, "OrchidBay_Hiring_2026_YTD.xlsx")]


def _lab8():
    e7 = _lab7()
    funnel = dict(applications=720, interviewed=sum(1 for r in YTD if r["Interview_Score"] is not None),
                  offered=sum(1 for r in YTD if r["Stage"] in ("Offered", "Hired")), hired=e7["hires"])
    cph = {s: v["cph"] for s, v in e7["src"].items()}
    cheapest = min((s for s in cph if cph[s]), key=lambda s: cph[s])
    dearest = max((s for s in cph if cph[s]), key=lambda s: cph[s])
    return dict(e7, funnel=funnel, cph=cph, cheapest=cheapest, dearest=dearest,
                share_li=round(e7["src"]["LinkedIn"]["apps"] / 720 * 100, 1))


def build_lab8(folder):
    wb = _ytd_workbook("Hiring 2026 Year to Date - Charts",
                       "The same 2026 year-to-date applications as Lab 7. In Lab 8 you build charts with Gemini, "
                       "learn which charts update with the data, and write titles that state the finding.")
    return [_save(wb, folder, "OrchidBay_Hiring_Charts.xlsx")]


# ---- dashboard (Labs 9-10)
def _dash_counts(pos="All", src="All"):
    rows = [r for r in YTD if (pos == "All" or r["Position"] == pos) and (src == "All" or r["Source"] == src)]
    hires = [r for r in rows if r["Stage"] == "Hired"]
    apps = len(rows)
    return dict(apps=apps, interviews=sum(1 for r in rows if r["Interview_Score"] is not None),
                offers=sum(1 for r in rows if r["Stage"] in ("Offered", "Hired")), hires=len(hires),
                rate=round(len(hires) / apps * 100, 1) if apps else 0,
                score=round(sum(r["Screening_Score"] for r in rows) / apps, 1) if apps else 0,
                dth=round(sum(r["Days_To_Hire"] for r in hires) / len(hires), 1) if hires else None,
                by_month={m: (sum(1 for r in rows if r["Month"] == m), sum(1 for r in hires if r["Month"] == m))
                          for m in sorted(MONTH_N and {f"2026-{k:02d}" for k in MONTH_N})},
                by_source={s: sum(1 for r in hires if r["Source"] == s) for s in SOURCES})


KPI_LABELS = ["Applications", "Interviews", "Offers", "Hires", "Hire rate", "Avg screening score", "Avg days to hire"]


def _dash_formulas():
    P = 'IF($C$3="All","*",$C$3)'
    S = 'IF($C$4="All","*",$C$4)'
    base = f'Applications!$F$2:$F$721,{P},Applications!$G$2:$G$721,{S}'
    return [
        f"=COUNTIFS({base})",
        f'=COUNTIFS({base},Applications!$L$2:$L$721,"<>")',
        f'=COUNTIFS({base},Applications!$M$2:$M$721,"Offered")+COUNTIFS({base},Applications!$M$2:$M$721,"Hired")',
        f'=COUNTIFS({base},Applications!$M$2:$M$721,"Hired")',
        '=IFERROR(E7/B7,0)',
        f'=IFERROR(AVERAGEIFS(Applications!$K$2:$K$721,{base}),0)',
        f'=IFERROR(AVERAGEIFS(Applications!$P$2:$P$721,{base},Applications!$M$2:$M$721,"Hired"),"-")',
    ], base


def _lists_sheet(wb):
    ls = wb.create_sheet("Lists")
    ls["A1"], ls["B1"] = "Positions", "Sources"
    ls["A1"].font = ls["B1"].font = Font(bold=True)
    for i, v in enumerate(["All"] + POSITIONS, 2):
        ls.cell(row=i, column=1, value=v)
    for i, v in enumerate(["All"] + SOURCES, 2):
        ls.cell(row=i, column=2, value=v)
    ls.column_dimensions["A"].width = 28
    ls.column_dimensions["B"].width = 20


def _dash_skeleton(wb, filled):
    d = wb.create_sheet("Dashboard", 1)
    d["A1"] = "Orchid Bay Talent Partners - Hiring Dashboard 2026 YTD"
    d["A1"].font = Font(bold=True, size=16, color="1A5E3A")
    d["A2"] = "Data as at 2026-09-30 · fictional training data"
    d["A2"].font = Font(italic=True, color="666666")
    d["B3"], d["B4"] = "Position", "Source"
    d["B3"].font = d["B4"].font = Font(bold=True)
    d["C3"], d["C4"] = "All", "All"
    for c in ("C3", "C4"):
        d[c].fill = PatternFill("solid", fgColor="FFF4CC")
    for i, lab in enumerate(KPI_LABELS):
        cell = d.cell(row=6, column=2 + i, value=lab)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3F1E8")
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        d.column_dimensions[get_column_letter(2 + i)].width = 16
    d.row_dimensions[6].height = 32
    d["A6"], d["A7"] = "KPI", "Value"
    # funnel B10:C14
    d["B10"], d["C10"] = "Funnel stage", "Count"
    for i, lab in enumerate(["Applications", "Interviewed", "Offered", "Hired"], 11):
        d.cell(row=i, column=2, value=lab)
    # monthly E10:G19
    d["E10"], d["F10"], d["G10"] = "Month", "Applications", "Hires"
    for i, m in enumerate([f"2026-{k:02d}" for k in MONTH_N], 11):
        d.cell(row=i, column=5, value=m)
    # sources I10:J17
    d["I10"], d["J10"] = "Source", "Hires"
    for i, s in enumerate(SOURCES, 11):
        d.cell(row=i, column=9, value=s)
    for c in ("B10", "C10", "E10", "F10", "G10", "I10", "J10"):
        d[c].font = Font(bold=True)
        d[c].fill = PatternFill("solid", fgColor="E3F1E8")
    d.column_dimensions["A"].width = 8
    d.column_dimensions["I"].width = 20
    if filled:
        f, base = _dash_formulas()
        for i, fx in enumerate(f):
            d.cell(row=7, column=2 + i, value=fx)
        d["F7"].number_format = "0.0%"
        d["G7"].number_format = "0.0"
        d["H7"].number_format = "0.0"
        d["C11"], d["C12"], d["C13"], d["C14"] = "=B7", "=C7", "=D7", "=E7"
        for i in range(11, 20):
            d.cell(row=i, column=6, value=f"=COUNTIFS({base},Applications!$D$2:$D$721,E{i})")
            d.cell(row=i, column=7, value=f'=COUNTIFS({base},Applications!$D$2:$D$721,E{i},Applications!$M$2:$M$721,"Hired")')
        for i in range(11, 18):
            d.cell(row=i, column=10, value=f'=COUNTIFS(Applications!$F$2:$F$721,IF($C$3="All","*",$C$3),'
                                           f'Applications!$G$2:$G$721,I{i},Applications!$M$2:$M$721,"Hired")')
        for cell, rng_ in (("C3", "Lists!$A$2:$A$10"), ("C4", "Lists!$B$2:$B$9")):
            dv = DataValidation(type="list", formula1=rng_, allow_blank=False)
            d.add_data_validation(dv)
            dv.add(cell)
    d.freeze_panes = "A5"


def _lab9():
    allv = _dash_counts()
    return dict(all=allv, se=_dash_counts("Software Engineer"), da_li=_dash_counts("Data Analyst", "LinkedIn"),
                ref=_dash_counts("All", "Referral"))


def build_lab9(folder):
    def extra(wb):
        _dash_skeleton(wb, filled=False)
        _lists_sheet(wb)
    wb = _ytd_workbook("Hiring Dashboard - Build the KPIs",
                       "The 2026 year-to-date applications, plus a Dashboard sheet with an empty layout and a Lists "
                       "sheet for the dropdowns. In Lab 9 you add dropdown selectors and KPI formulas with Gemini.",
                       extra=extra, costs=False)
    return [_save(wb, folder, "OrchidBay_Dashboard_Build.xlsx")]


def _lab10():
    return dict(_lab9(), cs=_dash_counts("Customer Success Executive"), ops=_dash_counts("Operations Supervisor"),
                hr=_dash_counts("HR Executive"))


def build_lab10(folder):
    def extra(wb):
        _dash_skeleton(wb, filled=True)
        _lists_sheet(wb)
    wb = _ytd_workbook("Hiring Dashboard - Make It Interactive",
                       "The finished Lab 9 dashboard: dropdowns in C3 (Position) and C4 (Source) drive the KPI row and "
                       "the three tables. In Lab 10 you add linked charts, a metric switch, conditional formatting and "
                       "protection, and test that everything responds to the dropdowns.", extra=extra, costs=False)
    return [_save(wb, folder, "OrchidBay_Dashboard_Interactive.xlsx")]


# =========================================================================== LABS 11-12: finalists
FINAL_ROLES = ["Data Analyst", "Software Engineer", "Customer Success Executive", "Finance Associate"]


def _finalists():
    rng = random.Random(1112)
    out = []
    for pos in FINAL_ROLES:
        pool = [r for r in Q3 if r["Position"] == pos and r["Stage"] in ("Shortlisted", "Interviewed", "Offered")]
        pool += [r for r in Q3 if r["Position"] == pos and r["Stage"] in ("Screened",) and r not in pool]
        pool = sorted(pool, key=lambda r: -r["Screening_Score"])[:4]
        for r in pool:
            fi = round(min(4.9, max(2.8, r["Screening_Score"] / 22 + rng.gauss(0, 0.45))), 1)
            out.append(dict(r, Final_Interview=fi,
                            Weighted=round(0.5 * r["Screening_Score"] + 0.5 * fi * 20, 1)))
    return out


FINALISTS = _finalists()
OPENINGS = {p: POS_JOB[p][8] for p in FINAL_ROLES}


def _must_have(r):
    req = [s.strip() for s in POS_JOB[r["Position"]][5].split(";")]
    return [s for s in req if s not in r["Skills_List"]]


def _optimum(budget):
    groups = [[r for r in FINALISTS if r["Position"] == p] for p in FINAL_ROLES]
    best = []
    for combo in itertools.product(*[itertools.combinations(g, OPENINGS[p]) for g, p in zip(groups, FINAL_ROLES)]):
        picks = [r for c in combo for r in c]
        sal = sum(r["Expected_Salary"] for r in picks)
        if sal <= budget:
            best.append((round(sum(r["Weighted"] for r in picks), 1), -sal, picks))
    best.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return best


def _budget():
    groups = [[r for r in FINALISTS if r["Position"] == p] for p in FINAL_ROLES]
    top = [r for g in groups for r in sorted(g, key=lambda r: -r["Weighted"])[:2]]
    return int((sum(r["Expected_Salary"] for r in top) - 1200) // 100 * 100), top


BUDGET, TOP_UNCONSTRAINED = _budget()


def _lab11():
    gaps = {r["Candidate_ID"]: _must_have(r) for r in FINALISTS}
    over = [r["Candidate_ID"] for r in FINALISTS if r["Expected_Salary"] > POS_JOB[r["Position"]][7]]
    da = [r for r in FINALISTS if r["Position"] == "Data Analyst"]
    return dict(n=len(FINALISTS), roles=FINAL_ROLES, gaps={k: v for k, v in gaps.items() if v}, over_budget=over,
                da_ids=[r["Candidate_ID"] for r in da],
                da_top=max(da, key=lambda r: r["Weighted"])["Candidate_ID"])


def _lab12():
    opt = _optimum(BUDGET)
    score, neg_sal, picks = opt[0]
    unc_score = round(sum(r["Weighted"] for r in TOP_UNCONSTRAINED), 1)
    unc_sal = sum(r["Expected_Salary"] for r in TOP_UNCONSTRAINED)
    assert opt[0][0] > opt[1][0] or opt[0][1] != opt[1][1], "optimum must be unique"
    return dict(budget=BUDGET, openings=OPENINGS, total_openings=sum(OPENINGS.values()),
                picks=sorted(r["Candidate_ID"] for r in picks), score=score, salary=-neg_sal,
                unc_score=unc_score, unc_salary=unc_sal,
                swapped_out=sorted(r["Candidate_ID"] for r in TOP_UNCONSTRAINED if r not in picks),
                swapped_in=sorted(r["Candidate_ID"] for r in picks if r not in TOP_UNCONSTRAINED),
                runner_up=opt[1][0])


FIN_COLS = ["Candidate_ID", "Full_Name", "Position", "Years_Experience", "Key_Skills", "Expected_Salary",
            "Budget_Max", "Notice_Weeks", "Screening_Score", "Final_Interview_Score", "Weighted_Score",
            "Candidate_Summary"]


def _fin_sheet(wb, title="Finalists"):
    ws = wb.active
    ws.title = title
    rows = [[r["Candidate_ID"], r["Full_Name"], r["Position"], r["Years_Experience"], r["Key_Skills"],
             r["Expected_Salary"], POS_JOB[r["Position"]][7], r["Notice_Weeks"], r["Screening_Score"],
             r["Final_Interview"], r["Weighted"], r["Summary"]] for r in FINALISTS]
    _write_rows(ws, FIN_COLS, rows, formats={"Expected_Salary": '"S$"#,##0', "Budget_Max": '"S$"#,##0'},
                widths={"Full_Name": 22, "Position": 26, "Key_Skills": 38, "Candidate_Summary": 70,
                        "Final_Interview_Score": 20, "Weighted_Score": 16})
    return ws


_FIN_DICT = [
    ("Finalists (sheet)", "—", "16 final-round candidates: the top four for each of four roles."),
    ("Final_Interview_Score", "Number", "Final panel interview score, 1.0-5.0."),
    ("Weighted_Score", "Number", "0.5 x Screening_Score + 0.5 x (Final_Interview_Score x 20), out of 100."),
    ("Budget_Max", "Number", "Top of the client's monthly salary band, S$."),
    ("Candidate_Summary", "Text", "Free text from the application."),
]


def build_lab11(folder):
    wb = Workbook()
    _fin_sheet(wb)
    _prompt_log_sheet(wb)
    _about_sheet(wb, "Final-Round Candidates",
                 "The 16 finalists for Data Analyst, Software Engineer, Customer Success Executive and Finance Associate. "
                 "In Lab 11 you build a Candidate Screener Gem in the Gemini app, give it the job descriptions and "
                 "screening rubric, and use it from the Gemini side panel in Sheets.", _FIN_DICT)
    paths = [_save(wb, folder, "OrchidBay_Finalists.xlsx")]
    paths += _knowledge_docs(folder)
    return paths


def build_lab12(folder):
    wb = Workbook()
    _fin_sheet(wb)
    c = wb.create_sheet("Constraints")
    _write_rows(c, ["Position", "Openings"], [[p, OPENINGS[p]] for p in FINAL_ROLES], freeze=False)
    c["A8"], c["B8"] = "Total monthly salary budget (S$)", BUDGET
    c["B8"].number_format = '"S$"#,##0'
    c["A8"].font = Font(bold=True)
    c.column_dimensions["A"].width = 34
    _prompt_log_sheet(wb)
    _about_sheet(wb, "Hiring Decision",
                 "The 16 finalists and the client's constraints: openings per role and one total monthly salary budget "
                 "for all eight hires. In Lab 12 you use Gemini's optimisation solver to pick the best combination of "
                 "hires, write the recommendation and share the results safely.",
                 _FIN_DICT + [("Constraints (sheet)", "—", "Openings per role and the total monthly salary budget.")])
    return [_save(wb, folder, "OrchidBay_Hiring_Decision.xlsx")]


# ---- Gem knowledge files (Lab 11): two Word documents to upload to Google Drive
JD_TEXT = {
    "Data Analyst": ("Build weekly finance dashboards, clean data from the core banking system and answer ad-hoc "
                     "questions from business teams.", ["SQL", "Excel", "Tableau"], ["Python", "Power BI"]),
    "Software Engineer": ("Develop and maintain cloud services for Kaya's SaaS platform in a small agile team.",
                          ["Python", "JavaScript", "Cloud"], ["CI/CD", "Docker"]),
    "Customer Success Executive": ("Onboard new SaaS customers, run check-in calls and reduce churn.",
                                   ["CRM", "Communication", "SaaS"], ["Customer Service", "Presentation"]),
    "Finance Associate": ("Prepare month-end reconciliations and support the accounts payable team.",
                          ["Accounting", "Excel", "Reconciliation"], ["Google Workspace", "Problem Solving"]),
    "HR Executive": ("Run end-to-end recruitment and HR administration for Orchid Bay's own staff.",
                     ["HRIS", "Recruitment", "Employment Act"], ["Payroll"]),
    "Payroll Specialist": ("Process monthly payroll and CPF contributions for 400 staff.",
                           ["Payroll", "CPF", "Excel"], ["HRIS"]),
    "Account Manager": ("Grow a portfolio of B2B retail accounts and negotiate annual contracts.",
                        ["B2B Sales", "CRM", "Negotiation"], ["Presentation"]),
    "Operations Supervisor": ("Lead a warehouse shift team of 15 and improve pick accuracy.",
                              ["Warehouse", "Team Leadership", "Lean"], ["Problem Solving"]),
}

RUBRIC = [
    ("Must-have skills", "4 = all must-have skills; 2 = missing one; 0 = missing two or more."),
    ("Relevant experience", "4 = 3+ years above the minimum; 3 = 1-2 years above; 2 = exactly the minimum; 0 = below."),
    ("Interview performance", "4 = Final_Interview_Score 4.5+; 3 = 4.0-4.4; 2 = 3.5-3.9; 1 = 3.0-3.4; 0 = below 3.0."),
    ("Salary fit", "4 = within the budget band; 2 = up to 10% above Budget_Max; 0 = more than 10% above."),
    ("Availability", "4 = notice of 4 weeks or less; 2 = 8 weeks; 0 = 12 weeks or more."),
]


def _knowledge_docs(folder):
    from docx import Document
    from docx.shared import Pt
    out = os.path.join(folder, "data", "gem-knowledge")
    os.makedirs(out, exist_ok=True)
    d = Document()
    d.add_heading("Orchid Bay Talent Partners - Job Descriptions (Q3 2026 Hiring Drive)", 1)
    d.add_paragraph(FICTION_NOTE)
    for j in JOBS:
        desc, must, nice = JD_TEXT[j[1]]
        d.add_heading(f"{j[0]} - {j[1]} ({j[2]})", 2)
        d.add_paragraph(desc)
        d.add_paragraph(f"Minimum experience: {j[4]} year{'s' if j[4] != 1 else ''}.")
        d.add_paragraph("Must-have skills: " + ", ".join(must) + ".")
        d.add_paragraph("Nice-to-have skills: " + ", ".join(nice) + ".")
        d.add_paragraph(f"Monthly salary band: S${j[6]:,} - S${j[7]:,}. Openings: {j[8]}.")
    p1 = os.path.join(out, "OrchidBay_Job_Descriptions.docx")
    d.save(p1)
    r = Document()
    r.add_heading("Orchid Bay Talent Partners - Candidate Screening Rubric", 1)
    r.add_paragraph(FICTION_NOTE)
    r.add_paragraph("Score every candidate on five criteria, 0-4 each, for a total out of 20. Use only the evidence "
                    "in the candidate's row and the job description for their role.")
    t = r.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Criterion", "How to score (0-4)"
    for crit, how in RUBRIC:
        row = t.add_row().cells
        row[0].text, row[1].text = crit, how
    r.add_heading("Recommendation", 2)
    r.add_paragraph("Advance = 16-20 and no must-have gap. Hold = 11-15, or 16+ with one must-have gap. Decline = 10 or "
                    "less.")
    r.add_heading("Fair screening rules", 2)
    for rule in ["Never use or guess age, gender, race, religion, nationality, marital status or health.",
                 "Do not treat a name, school or photo as evidence of ability.",
                 "Say 'not enough information' rather than guess.",
                 "The recommendation supports the hiring manager - a person makes the decision."]:
        r.add_paragraph(rule, style="List Bullet")
    for para in r.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)
    p2 = os.path.join(out, "OrchidBay_Screening_Rubric.docx")
    r.save(p2)
    return [p1, p2]


_BUILDERS = {1: build_lab1, 2: build_lab2, 3: build_lab3, 4: build_lab4, 5: build_lab5, 6: build_lab6,
             7: build_lab7, 8: build_lab8, 9: build_lab9, 10: build_lab10, 11: build_lab11, 12: build_lab12}

EXPECTED = {1: _lab1(), 2: _lab2(), 3: _lab3(), 4: _lab4(), 5: _lab5(), 6: _lab6(), 7: _lab7(), 8: _lab8(),
            9: _lab9(), 10: _lab10(), 11: _lab11(), 12: _lab12()}


def build_lab_data(num, folder):
    return _BUILDERS[num](folder) if num in _BUILDERS else []


if __name__ == "__main__":
    import json
    print(json.dumps(EXPECTED, indent=1, default=str))
